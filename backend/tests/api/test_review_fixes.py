"""Regression tests from the whole-branch review: PUT re-validation, short-term
rejection, input completeness, non-PSD correlation, run-snapshot export."""
from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import create_app

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "golden_excelcior.json"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SUPABASE_URL", "")
    monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "")
    from app.core.settings import get_settings
    get_settings.cache_clear()
    from app.repositories.store import MemoryStore
    app = create_app()
    app.state.store = MemoryStore(persist_path=str(tmp_path / "store.json"))
    with TestClient(app) as c:
        yield c


def _fx() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def _new_tx(client, payload: dict) -> dict:
    batch = client.post("/api/batches", json={"name": "b"}).json()
    return client.post(f"/api/batches/{batch['id']}/transactions", json=payload).json()


def test_put_then_calculate_with_portfolio_and_dcr_changes(client):
    fx = _fx()
    tx = _new_tx(client, fx["inputs"])
    sheet = fx["portfolio_sheet"]
    edited = fx["inputs"] | {
        "portfolio_of_shares": True,
        "portfolio": {"stocks": sheet["stocks"], "equity_correlation": sheet["equity_correlation"]},
        "dcr_changes": True,
        "dcr_schedule": [{"time": 3, "dcr": 1.5}],
        "rate_linked_to_cover": False,
    }
    r = client.put(f"/api/transactions/{tx['id']}", json=edited)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "draft" and r.json()["latest_pd"] is None
    r = client.post(f"/api/transactions/{tx['id']}/calculate")
    assert r.status_code == 200, r.text
    run = r.json()
    assert run["result"]["derived"]["dcr_by_year"][3] == 1.5
    r = client.get(f"/api/runs/{run['id']}/export.xlsx")
    assert r.status_code == 200, r.text


def test_term_of_six_months_or_less_is_rejected(client):
    tx = _new_tx(client, _fx()["inputs"] | {"maturity_date": "2015-11-30"})
    r = client.post(f"/api/transactions/{tx['id']}/calculate")
    assert r.status_code == 422 and "six months" in r.text.lower()


@pytest.mark.parametrize(
    "patch, needle",
    [
        ({"structure": "Straight Amortisation", "amortised_loan_amount": 5e7, "bullet_loan_amount": None,
          "amortised_interest_rate": None}, "interest rate charged"),
        ({"bullet_interest_rate": None}, "interest rate on bullet"),
        ({"average_daily_volume": 0}, "daily trading volume"),
    ],
)
def test_missing_engine_inputs_are_rejected(client, patch, needle):
    tx = _new_tx(client, _fx()["inputs"] | patch)
    r = client.post(f"/api/transactions/{tx['id']}/calculate")
    assert r.status_code == 422 and needle in r.text.lower(), r.text


def test_empty_portfolio_value_is_rejected(client):
    tx = _new_tx(client, _fx()["inputs"] | {"portfolio_of_shares": True, "portfolio": {"stocks": [{"name": "x"}]}})
    r = client.post(f"/api/transactions/{tx['id']}/calculate")
    assert r.status_code == 422 and "portfolio" in r.text.lower()


def test_non_psd_correlation_is_rejected(client):
    stocks = [{"name": n, "number_of_shares": 100, "share_price": 10, "volatility": 0.3,
               "country": "South Africa", "risk_grade": "SB10"} for n in "abc"]
    corr = [[1.0 if i == j else (-1.0 if i < 3 and j < 3 else 0.0) for j in range(10)] for i in range(10)]
    tx = _new_tx(client, _fx()["inputs"] | {"portfolio_of_shares": True,
                                            "portfolio": {"stocks": stocks, "equity_correlation": corr}})
    r = client.post(f"/api/transactions/{tx['id']}/calculate")
    assert r.status_code == 422 and "correlation" in r.text.lower()


def test_export_uses_the_run_snapshot(client):
    tx = _new_tx(client, _fx()["inputs"])
    run = client.post(f"/api/transactions/{tx['id']}/calculate").json()
    client.put(f"/api/transactions/{tx['id']}", json=_fx()["inputs"] | {"client_name": "Renamed Co"})
    r = client.get(f"/api/runs/{run['id']}/export.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Share Cover Risk Report"]
    assert "Excelcior Shipping" in str(ws["A2"].value) and "Renamed" not in str(ws["A2"].value)
