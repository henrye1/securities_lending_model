import io
import json
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import create_app

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "golden_excelcior.json"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SUPABASE_URL", "")
    monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "")
    from app.core.settings import get_settings
    get_settings.cache_clear()
    from app.repositories.store import MemoryStore
    app = create_app()
    app.state.store = MemoryStore(persist_path=str(tmp_path / "store.json"))
    with TestClient(app) as c:
        yield c


def test_export_xlsx(client):
    inputs = json.loads(FIXTURE.read_text(encoding="utf-8"))["inputs"]
    batch = client.post("/api/batches", json={"name": "x"}).json()
    tx = client.post(f"/api/batches/{batch['id']}/transactions", json=inputs).json()
    run = client.post(f"/api/transactions/{tx['id']}/calculate").json()
    r = client.get(f"/api/runs/{run['id']}/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "12345678_Excelcior-Shipping_" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Share Cover Risk Report", "Calculations", "Inputs"]
    ws = wb["Share Cover Risk Report"]
    assert str(ws["A1"].value).startswith("CYTE Risk Suite")
    assert str(ws["A2"].value).startswith("SHARE COVER RISK REPORT")
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "SB19" in values and "Combined" in values and "Bullet Payment" in values
