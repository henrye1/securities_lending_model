"""Excel export of a run - CYTE Risk Suite branding.

Sheets: Share Cover Risk Report (the workbook's Print Report & Results layout)
| Calculations (year table + results block) | Inputs (the Input page).
Filename: {CIF}_{Client}_{YYMMDD}.xlsx"""
from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.models import Run, Transaction

NAVY = "001080"
PANEL = "DBE3F1"
BORDER = Border(*(Side(style="thin", color="B0B8CC"),) * 4)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SUB_FONT = Font(bold=True, size=11, color=NAVY)
HEAD_FONT = Font(bold=True, size=10, color=NAVY)
HEAD_FILL = PatternFill("solid", start_color=PANEL)
BRAND_FONT = Font(bold=True, size=10, color="FFFFFF")
BRAND_FILL = PatternFill("solid", start_color=NAVY)

PCT, PCT4, NUM, RATIO, DATE = "0.00%", "0.0000%", "#,##0.00", "0.00", "yyyy-mm-dd"

YEAR_ROWS = (
    ("dt", "Time Change", RATIO),
    ("dcr", "Default Cover Ratio", RATIO),
    ("avg_loan", "Average Loan Amount", NUM),
    ("avg_share_value", "Average Share Value", NUM),
    ("spv_pd_covenant", "SPV Annual PD [Covenant Breach]", PCT4),
    ("spv_pd", "SPV Annual PD [SPV]", PCT4),
    ("spv_pd_no_covenant", "SPV Annual PD [No Covenant]", PCT4),
    ("issuer_default_rate", "Issuer Default Rate", PCT4),
    ("issuer_lgd", "Issuer Loss Given Default", PCT),
    ("spv_lgd", "SPV Loss Given Default", PCT),
    ("covenant_effectiveness", "Covenant Effectiveness Index", RATIO),
    ("combined_pd", "Combined Risk", PCT4),
    ("combined_lgd", "Combined LGD", PCT),
)

RESULT_ROWS = (
    ("pd_client_avg", "PD Client - Average PD", PCT4),
    ("pd_client_weighted", "PD Client - Weighted PD", PCT4),
    ("pd_client_grade", "PD Client - Risk Grade", None),
    ("pd_issuer_avg", "PD Issuer - Average PD", PCT4),
    ("pd_issuer_weighted", "PD Issuer - Weighted PD", PCT4),
    ("pd_total", "PD - Total Transaction", PCT4),
    ("lgd_client_avg", "LGD Client - Average LGD", PCT),
    ("lgd_client_weighted", "LGD Client - Weighted LGD", PCT),
    ("lgd_issuer_avg", "LGD Issuer - Average LGD", PCT),
    ("lgd_issuer_weighted", "LGD Issuer - Weighted LGD", PCT),
    ("lgd_total", "LGD - Total Transaction", PCT),
    ("el", "EL", PCT4),
    ("risk_grade", "RG", None),
    ("rwa_pct", "RWA%", PCT),
)

REPORT_ROWS = (
    ("pd", "PD", PCT4),
    ("grade", "Risk Grade", None),
    ("lgd", "Downturn LGD", PCT),
    ("el", "EL", PCT4),
    ("el_grade", "EL Grade", None),
    ("rwa_pct", "RWA%", PCT),
)


def _brand_band(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = BRAND_FILL
        c.font = BRAND_FONT
    ws.cell(row=1, column=1, value="CYTE Risk Suite  |  Securities Lending (Share Cover) Model")


def _sheet_title(ws, title: str, ncols: int, widths: tuple[int, ...] = ()) -> None:
    _brand_band(ws, ncols)
    ws.cell(row=2, column=1, value=title).font = TITLE_FONT
    for col in range(1, ncols + 1):
        w = widths[col - 1] if col - 1 < len(widths) else 18
        ws.column_dimensions[get_column_letter(col)].width = w


def _header_row(ws, row: int, values: list) -> None:
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center" if col > 1 else "left")


def _cell(ws, row: int, col: int, value, fmt: str | None = None, bold: bool = False) -> None:
    c = ws.cell(row=row, column=col, value="" if value is None else value)
    c.border = BORDER
    if fmt and isinstance(value, (int, float)):
        c.number_format = fmt
    if bold:
        c.font = Font(bold=True)
    if col > 1:
        c.alignment = Alignment(horizontal="right" if isinstance(value, (int, float)) else "center")


def _na(v, fmt: str | None = None):
    return "Not Applicable" if v is None else v


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-") or "transaction"


def export_run(tx: Transaction, run: Run) -> tuple[str, bytes]:
    res = run.result
    derived, report, results, years = res["derived"], res["report"], res["results"], res["years"]
    wb = Workbook()

    # ---- Share Cover Risk Report ------------------------------------------
    ws = wb.active
    ws.title = "Share Cover Risk Report"
    _sheet_title(ws, f"SHARE COVER RISK REPORT FOR: {tx.client_name}", 4, (44, 22, 22, 22))
    r = 4
    header = (
        ("CIF Code of Client / Borrower", tx.client_cif),
        ("Date", res.get("report_date", run.created_at.date().isoformat())),
        ("Reference Share", tx.reference_name),
        ("CIF Code of Reference Entity", tx.reference_cif),
    )
    for label, value in header:
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, value)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Transaction Details").font = SUB_FONT
    r += 1
    details = (
        ("Start Date", tx.start_date.isoformat(), None),
        ("Maturity Date", tx.maturity_date.isoformat(), None),
        ("Years", derived["years_total"], RATIO),
        ("Amortised Loan Amount" if tx.structure != "Interest Only Repayments" else "Interest Only Loan Amount",
         derived["amortised_amount"], NUM),
        ("Bullet Loan Amount", derived["bullet_amount"], NUM),
        ("Interest Rate on Bullet Loan", derived["bullet_rate"], PCT),
        ("Number of Shares", derived["number_of_shares"], NUM),
        ("Share Price", derived["share_price"], "#,##0.0000"),
        ("Initial Share Value", derived["share_value"], NUM),
        ("Initial Cover Ratio", derived["initial_cover_ratio"], RATIO),
        ("Initial Default Cover Ratio", derived["initial_dcr"], RATIO),
        ("PD of Reference Entity", derived["reference_pd"], PCT4),
        ("Dividend Yield", derived["dividend_yield"], PCT),
        ("Percentage of Dividends Repayed", tx.dividends_repaid, PCT),
        ("Share Price Volatility", derived["volatility"], PCT),
        ("Liabilities Ranking Ahead", tx.liabilities_ahead or 0, NUM),
        ("Liabilities Ranking Parri-Passu", tx.liabilities_pari_passu or 0, NUM),
        ("Liabilities Ranking Subordinate", tx.liabilities_subordinate or 0, NUM),
        ("Transaction Structure", tx.structure, None),
    )
    for label, value, fmt in details:
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, value, fmt)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Results").font = SUB_FONT
    r += 1
    _header_row(ws, r, ["", "Reference Entity", "Client/Borrower", "Combined"])
    r += 1
    for key, label, fmt in REPORT_ROWS:
        _cell(ws, r, 1, label, bold=True)
        for col, block in enumerate(("reference", "client", "combined"), start=2):
            _cell(ws, r, col, report[block][key], fmt)
        r += 1
    if res.get("warnings"):
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = SUB_FONT
        for w in res["warnings"]:
            r += 1
            ws.cell(row=r, column=1, value=w)

    # ---- Calculations -------------------------------------------------------
    ws = wb.create_sheet("Calculations")
    n = len(years)
    _sheet_title(ws, f"STOCK LENDING MODEL - {tx.client_name}", 1 + n, (36,) + (14,) * n)
    _header_row(ws, 4, ["Year"] + [y["year"] for y in years])
    r = 5
    for key, label, fmt in YEAR_ROWS:
        _cell(ws, r, 1, label, bold=True)
        for col, y in enumerate(years, start=2):
            _cell(ws, r, col, y[key], fmt)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Results").font = SUB_FONT
    r += 1
    for key, label, fmt in RESULT_ROWS:
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, results[key], fmt)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Calculation Factors").font = SUB_FONT
    r += 1
    factors = (
        ("Recovery Policy", derived["recovery_policy"], None),
        ("Listing", tx.listing, None),
        ("Liquidity", derived["liquidity"], None),
        ("Fire-sale haircut (FSH)", derived["fsh"], PCT),
        ("Time to unwind (years)", derived["unwind_years"], "0.0000"),
        ("Time to unwind (days)", derived["unwind_days"], "0.0"),
        ("Discount Factor", derived["discount_factor"], "0.000000"),
        ("Risk Free Rate", derived["risk_free_rate"], PCT),
        ("Random draws", "Workbook draw set" if derived["workbook_draws"] else f"Seed {tx.random_seed}", None),
        ("Simulation paths", derived["n_paths"], "#,##0"),
    )
    for label, value, fmt in factors:
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, value, fmt)
        r += 1

    # ---- Inputs ---------------------------------------------------------------
    ws = wb.create_sheet("Inputs")
    _sheet_title(ws, "INPUT PAGE", 2, (48, 30))
    r = 4
    inputs = (
        ("Client / Borrower Details", None),
        ("CIF Code of Client / Borrower", tx.client_cif),
        ("Client / Borrowing Entity Name", tx.client_name),
        ("Country of Risk", tx.client_country),
        ("Portfolio of Shares?", "Yes" if tx.portfolio_of_shares else "No"),
        ("Transaction Details", None),
        ("Start Date", tx.start_date.isoformat()),
        ("Maturity Date", tx.maturity_date.isoformat()),
        ("Transaction Structure", tx.structure),
        ("Amortised / Interest Only Loan Amount", _na(tx.amortised_loan_amount)),
        ("Bullet Loan Amount", _na(tx.bullet_loan_amount)),
        ("Number of Shares", derived["number_of_shares"]),
        ("Share Price", derived["share_price"]),
        ("Total Stock Value", derived["share_value"]),
        ("Initial Cover Ratio", derived["initial_cover_ratio"]),
        ("Average Daily Trading Volumes", tx.average_daily_volume),
        ("Does Default Cover Ratio Change?", "Yes" if tx.dcr_changes else "No"),
        ("Default Cover Ratio", tx.default_cover_ratio),
        ("DCR changes (time: dcr)", "; ".join(f"{c.time}: {c.dcr}" for c in tx.dcr_schedule) or "-"),
        ("Interest", None),
        ("Current Interest Rate on Bullet Loan", _na(tx.bullet_interest_rate)),
        ("Is Interest Rate Linked to Cover Ratio?", "Yes" if tx.rate_linked_to_cover else "No"),
        ("Current Interest Rate Charged (amortised loan)", _na(tx.amortised_interest_rate)),
        ("Cover-linked rates (lower bound: rate)",
         "; ".join(f"{b.lower_bound}: {b.rate}" for b in tx.rate_schedule) or "-"),
        ("Stock Details", None),
        ("Dividend Yield", tx.dividend_yield),
        ("Percentage of Dividends Paid Back", tx.dividends_repaid),
        ("Stock Volatility", tx.stock_volatility),
        ("Time to unwind share in market (years)", derived["unwind_years"]),
        ("Other Liabilities and Collateral", None),
        ("Liabilities Ranking Ahead of Facility", _na(tx.liabilities_ahead)),
        ("Liabilities Ranking Parri-Passu with Facility", _na(tx.liabilities_pari_passu)),
        ("Liabilities Subordinate to Facility", _na(tx.liabilities_subordinate)),
        ("Other Collateral Pledged to Facility", _na(tx.other_collateral)),
        ("Collateral Pledged to Parri-Passu Liabilities", _na(tx.collateral_pari_passu)),
        ("Other Parri-Passu Liabilities to Which Collateral is Pledged", _na(tx.other_pari_passu_liabilities)),
        ("Reference Share Details", None),
        ("CIF Code of Reference Entity", tx.reference_cif),
        ("Reference Entity / Share Name", tx.reference_name),
        ("Country of Risk", tx.reference_country),
        ("Reference Entity Share Listing", tx.listing),
        ("Risk Grade of Reference Entity", tx.risk_grade),
        ("Revenue Size", tx.revenue_size),
        ("Corporate Sector(s)", "; ".join(f"{s.sector} {s.weight:.0%}" for s in tx.sectors) or "-"),
        ("Unlisted Share Haircut", tx.unlisted_haircut),
        ("Illiquid Share Haircut", tx.illiquid_haircut),
        ("Risk Free Rate", derived["risk_free_rate"]),
        ("Currency Risk", None),
        ("Is the equity and loan in different currencies?", "Yes" if tx.currency_mismatch else "No"),
        ("Currency (loan/equity) volatility", tx.currency_volatility),
        ("Correlation between currency and equity", tx.currency_correlation),
        ("Calculation Settings", None),
        ("Recovery Policy", derived["recovery_policy"]),
        ("Random Seed", "Workbook draw set" if tx.random_seed is None else tx.random_seed),
    )
    for label, value in inputs:
        if value is None:
            r += 1
            ws.cell(row=r, column=1, value=label).font = SUB_FONT
        else:
            _cell(ws, r, 1, label, bold=True)
            _cell(ws, r, 2, value)
        r += 1
    port = res.get("portfolio")
    if port and tx.portfolio:
        r += 1
        ws.cell(row=r, column=1, value="Portfolio Transactions - Outputs (Combined)").font = SUB_FONT
        for key in ("number_of_shares", "share_price", "stock_volatility", "dividend_yield",
                    "average_daily_volume", "time_to_unwind", "unlisted_haircut", "illiquid_haircut",
                    "risk_free_rate", "equity_lgd"):
            r += 1
            _cell(ws, r, 1, key.replace("_", " ").title(), bold=True)
            _cell(ws, r, 2, port.get(key))

    stamp = datetime.now().strftime("%y%m%d")
    filename = f"{tx.client_cif or 'NOCIF'}_{_slug(tx.client_name)}_{stamp}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    return filename, buf.getvalue()
